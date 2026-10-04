"""lab_steps.py - phần "keo" cho notebook: tổng hợp kết quả, suy luận Bước 3, chung kết Bước 4, xlsx Bước 5.

Mọi quyết định ở đây chỉ dùng VAL. Hàm duy nhất chạm test là final_predictions(), gọi ở Bước 4,
mỗi seed một lần, với phương pháp suy luận đã chốt trên val từ trước.
"""
from __future__ import annotations

import dataclasses
import json
from pathlib import Path

import numpy as np
import pandas as pd

import inference as inf
from dataset import CLASS_NAMES, build_transforms, load_split, make_loader
from train import Config, load_trained, run_dir

# --------------------------------------------------------------------------- #
# Đọc kết quả các lần chạy
# --------------------------------------------------------------------------- #

def load_summaries(out_dir: str | Path) -> pd.DataFrame:
    rows = []
    for f in sorted(Path(out_dir).glob("*/seed*/summary.json")):
        s = json.loads(f.read_text())
        rows.append(s)
    return pd.DataFrame(rows)


def summary_of(out_dir, exp_id: str, seed: int = 0) -> dict | None:
    f = Path(out_dir) / exp_id / f"seed{seed}" / "summary.json"
    return json.loads(f.read_text()) if f.exists() else None


def cfg_of(out_dir, exp_id: str, seed: int = 0) -> Config:
    d = json.loads((Path(out_dir) / exp_id / f"seed{seed}" / "config.json").read_text())
    return Config(**d)


def clone_run(out_dir, src: str, dst: str, seed: int = 0, curves_dir: str | None = None) -> None:
    """T00 seed0 ≡ B0x seed0 (cùng công thức nền, cùng seed): chép kết quả thay vì train lại.
    Ghi rõ trong báo cáo. config.json/summary.json được đổi exp_id thành dst."""
    import shutil

    s, d = Path(out_dir) / src / f"seed{seed}", Path(out_dir) / dst / f"seed{seed}"
    if (d / "summary.json").exists():
        return
    d.parent.mkdir(parents=True, exist_ok=True)
    shutil.copytree(s, d, dirs_exist_ok=True)
    for name in ("config.json", "summary.json"):
        j = json.loads((d / name).read_text())
        j["exp_id"] = dst
        if "config" in j:
            j["config"]["exp_id"] = dst
        if name == "summary.json":
            j["cloned_from"] = src
            if curves_dir and j.get("curve"):
                src_curve = Path(j["curve"])
                new_curve = Path(curves_dir) / src_curve.name.replace(src, dst, 1)
                if src_curve.exists():
                    shutil.copy(src_curve, new_curve)
                j["curve"] = str(new_curve)
        (d / name).write_text(json.dumps(j, indent=2, default=float))


def backbone_table(df: pd.DataFrame) -> pd.DataFrame:
    b = df[df.exp_id.str.startswith("B")].sort_values("exp_id")
    return pd.DataFrame({
        "exp_id": b.exp_id, "backbone": b.backbone, "weight_tag": b.weight_tag,
        "params_M": b.params_M.round(2), "GMAC": b.gmacs.round(3), "img_size": b.img_size,
        "epochs": b.epochs, "seed": b.seed, "best_epoch": b.best_epoch,
        "val_macro_f1": b.val_macro_f1, "val_top1": b.val_top1,
        "train_time_per_epoch_s": b.train_time_per_epoch_s.round(1),
        "latency_b1_p50_ms": b.latency_b1_p50_ms, "latency_b1_p95_ms": b.latency_b1_p95_ms,
        "curve": b.curve,
    }).reset_index(drop=True)


BASE = Config()


def diff_from_base(cfg: dict, base: dict) -> str:
    skip = {"exp_id", "seed", "desc", "pred_tag", "out_dir", "pred_dir", "curves_dir", "images_dir",
            "labels_dir", "measure_latency", "skip_if_done", "save_checkpoint", "save_test_predictions",
            "num_workers", "max_steps_per_epoch"}
    d = [f"{k}={cfg[k]}" for k in cfg if k not in skip and cfg.get(k) != base.get(k)]
    return ", ".join(d) if d else "(nền T00)"


def noise_std(df: pd.DataFrame, exp_id: str = "T00") -> tuple[float, float, int]:
    v = df[df.exp_id == exp_id].val_macro_f1.values
    return (float(v.mean()), float(v.std(ddof=1)) if len(v) > 1 else float("nan"), len(v))


def training_table(df: pd.DataFrame, axes: dict[str, str], base_id: str = "T00") -> pd.DataFrame:
    base = df[(df.exp_id == base_id) & (df.seed == 0)].iloc[0]
    _, std, nseed = noise_std(df, base_id)
    rows = []
    t = df[df.exp_id.str.startswith("T")].sort_values(["exp_id", "seed"])
    for _, r in t.iterrows():
        delta = r.val_macro_f1 - base.val_macro_f1
        f1c = r.val_f1_per_class
        verdict = ""
        if r.exp_id != base_id and not np.isnan(std):
            verdict = ("tốt hơn rõ (Δ > 2·std)" if delta > 2 * std else
                       "kém hơn rõ (Δ < -2·std)" if delta < -2 * std else "không phân biệt được (|Δ| ≤ 2·std)")
        rows.append({
            "exp_id": r.exp_id, "backbone": r.backbone, "axis": axes.get(r.exp_id, "nền"),
            "diff_vs_T00": diff_from_base(r.config, base.config), "seed": r.seed,
            "val_macro_f1": r.val_macro_f1, "val_top1": r.val_top1,
            "delta_vs_T00": delta, "noise_std_T00": std, "verdict": verdict,
            "f1_chinee_apple": f1c[0], "f1_snake_weed": f1c[7], "best_epoch": r.best_epoch,
            "curve": r.curve,
        })
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------- #
# Bước 3: suy luận trên VAL
# --------------------------------------------------------------------------- #

METHODS = {
    # tên: (loại loader, kích thước, hàm sinh view, không gian gộp, K)
    "I00_1view":        ("std", 224, lambda x: [x], "prob", 1),
    "I01_hflip_prob":   ("std", 224, inf.views_hflip_pair, "prob", 2),
    "I03_hflip_logit":  ("std", 224, inf.views_hflip_pair, "logit", 2),
    "I02_5crop_prob":   ("full", 256, lambda x: inf.views_multicrop(x, 224), "prob", 5),
    "I02_10crop_prob":  ("full", 256, lambda x: inf.views_multicrop(x, 224, flip=True), "prob", 10),
}


def _loader(df, images_dir, kind: str, size: int, batch_size=64, num_workers=2):
    tf = build_transforms(False, size) if kind == "std" else inf.eval_transform_full(size)
    return make_loader(df, images_dir, tf, batch_size, train=False, num_workers=num_workers)


def run_method(model, df, images_dir, method: str, device, amp=False, num_workers=2):
    kind, size, vf, space, k = METHODS[method]
    loader = _loader(df, images_dir, kind, size, num_workers=num_workers)
    names, y, logits_list = inf.predict_views(model, loader, device, vf, amp=amp)
    return names, y, logits_list, space


def run_resolution(model, df, images_dir, size: int, device, num_workers=2):
    loader = _loader(df, images_dir, "std", size, num_workers=num_workers)
    names, y, logits = inf.predict_logits(model, loader, device)
    return names, y, [logits], "prob"


def calibrated_probs(logits_list, T: float = 1.0, space: str = "prob"):
    if space == "logit":
        return inf.apply_temperature(np.mean(np.stack(logits_list), 0), T)
    return inf.ensemble_probs([inf.apply_temperature(l, T) for l in logits_list])


def fit_T(logits_list, y, space: str = "prob") -> float:
    """T cực tiểu NLL của xác suất cuối (sau gộp view) trên VAL."""
    from scipy.optimize import minimize_scalar

    def nll(logt):
        p = calibrated_probs(logits_list, float(np.exp(logt)), space)
        return -np.log(np.clip(p[np.arange(len(y)), y], 1e-12, None)).mean()

    r = minimize_scalar(nll, bounds=(np.log(0.05), np.log(20.0)), method="bounded")
    return float(np.exp(r.x))


def metrics(ev, y, probs) -> dict:
    m = ev.compute_metrics(np.asarray(y), probs.argmax(1), probs)
    return {k: m[k] for k in ("macro_f1", "top1", "ece", "nll", "balanced_acc")} | {"f1": m["f1"], "recall": m["recall"]}


def crossfit_ece(ev, logits_list, y, space="prob", seed=0) -> tuple[float, float]:
    """ECE 'trung thực' trên val: khớp T trên một nửa val, đo trên nửa còn lại (2 chiều, lấy trung bình)."""
    rng = np.random.default_rng(seed)
    idx = rng.permutation(len(y))
    a, b = idx[: len(y) // 2], idx[len(y) // 2:]
    before, after = [], []
    for fit, test in ((a, b), (b, a)):
        T = fit_T([l[fit] for l in logits_list], y[fit], space)
        before.append(ev.ece_score(calibrated_probs([l[test] for l in logits_list], 1.0, space), y[test]))
        after.append(ev.ece_score(calibrated_probs([l[test] for l in logits_list], T, space), y[test]))
    return float(np.mean(before)), float(np.mean(after))


# --------------------------------------------------------------------------- #
# Bước 4: chung kết (TEST một lần mỗi seed)
# --------------------------------------------------------------------------- #

def final_predictions(ev, cfg: Config, method: str, images_dir, labels_dir, pred_dir, device,
                      tag: str = "F01", num_workers: int = 2) -> dict:
    """Áp dụng phương pháp suy luận ĐÃ CHỐT TRÊN VAL cho checkpoint của cfg:
       - khớp T trên VAL (sau gộp view), ghi <tag>_seed<k>_val.csv (đã hiệu chuẩn)
       - chạy TEST đúng một lần: ghi <tag>_seed<k>_test.csv (đã hiệu chuẩn) và <tag>uncal_seed<k>_test.csv
    """
    train_df, val_df, test_df = load_split(labels_dir, cfg.fold)
    model = load_trained(cfg, device)
    res_size = int(method.split("res")[1]) if method.startswith("I04_res") else None

    def go(df):
        if res_size:
            return run_resolution(model, df, images_dir, res_size, device, num_workers)
        return run_method(model, df, images_dir, method, device, num_workers=num_workers)

    vn, vy, vl, space = go(val_df)
    T = fit_T(vl, vy, space)
    pv = calibrated_probs(vl, T, space)
    k = cfg.seed
    ev.save_predictions(Path(pred_dir) / f"{tag}_seed{k}_val.csv", vn, vy, pv)
    tn, ty, tl, _ = go(test_df)  # TEST: một lần
    rd = run_dir(cfg)
    np.save(rd / f"{tag}_{method}_test_logits.npy", np.stack(tl))
    p_unc = calibrated_probs(tl, 1.0, space)
    p_cal = calibrated_probs(tl, T, space)
    ev.save_predictions(Path(pred_dir) / f"{tag}uncal_seed{k}_test.csv", tn, ty, p_unc)
    ev.save_predictions(Path(pred_dir) / f"{tag}_seed{k}_test.csv", tn, ty, p_cal)
    mv, mt, mu = metrics(ev, vy, pv), metrics(ev, ty, p_cal), metrics(ev, ty, p_unc)
    out = {"seed": k, "method": method, "T": T, "val_macro_f1": mv["macro_f1"], "val_top1": mv["top1"],
           "test_macro_f1": mt["macro_f1"], "test_top1": mt["top1"], "test_ece": mt["ece"],
           "test_ece_uncal": mu["ece"], "test_balanced_acc": mt["balanced_acc"]}
    (rd / f"{tag}_final.json").write_text(json.dumps(out, indent=2))
    return out


# --------------------------------------------------------------------------- #
# Bước 5: results.xlsx
# --------------------------------------------------------------------------- #

def write_xlsx(path, sheets: dict[str, pd.DataFrame], highlight: dict[str, str] | None = None) -> None:
    """Ghi các sheet, cố định hàng tiêu đề, 4 chữ số thập phân, tô dòng tốt nhất theo cột `highlight[sheet]`."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    highlight = highlight or {}
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df = df.copy()
            for c in df.columns:
                if df[c].dtype == object:
                    df[c] = df[c].apply(lambda v: ", ".join(f"{x:.4f}" if isinstance(x, float) else str(x) for x in v)
                                        if isinstance(v, (list, tuple, np.ndarray)) else v)
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for j, col in enumerate(df.columns, 1):
                longest = df[col].astype(str).str.len().max() if len(df) else 10
                longest = 10 if pd.isna(longest) else int(longest)
                width = max(10, min(45, longest + 2), len(str(col)) + 2)
                ws.column_dimensions[get_column_letter(j)].width = width
                if pd.api.types.is_float_dtype(df[col]):
                    for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                        for cell in row:
                            cell.number_format = "0.0000"
            hc = highlight.get(name)
            if hc and hc in df.columns and len(df) and df[hc].notna().any():
                best_row = int(pd.to_numeric(df[hc], errors="coerce").idxmax()) + 2
                for cell in ws[best_row]:
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
